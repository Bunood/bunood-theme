"""Actual native report exports and transactional compatibility acceptance."""

from io import BytesIO
import json
from pathlib import Path
import uuid
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase
from openpyxl import load_workbook

from bunood_theme import report_compat


class TestReportCompatibility(IntegrationTestCase):
    def setUp(self):
        super().setUp()
        self.previous_user = frappe.session.user
        frappe.set_user("Administrator")
        self.savepoint = "reports_" + uuid.uuid4().hex
        frappe.db.savepoint(self.savepoint)

    def tearDown(self):
        frappe.set_user("Administrator")
        frappe.db.rollback(save_point=self.savepoint)
        frappe.set_user(self.previous_user)
        super().tearDown()

    def owned_query(self):
        name = "BND Report acceptance " + uuid.uuid4().hex[:12]
        frappe.get_doc(dict(doctype="Report", report_name=name, ref_doctype="ToDo", module="Core",
                            report_type="Query Report", is_standard="No",
                            query='SELECT 15 AS "Probe:Int:100"', roles=[dict(role="System Manager")])).insert()
        # Only change the owned row: standard .insert() would export to source.
        frappe.db.set_value("Report", name, "is_standard", "Yes")
        frappe.clear_document_cache("Report", name)
        return name

    def params(self, name):
        return frappe._dict(report_name=name, file_format_type="Excel", filters={}, ignore_visible_idx=1)

    def test_sql_only_native_xlsx_export_has_real_value(self):
        from frappe.desk.query_report import _export_query
        name = self.owned_query()
        report_name, extension, content = _export_query(self.params(name), frappe._dict(), populate_response=False)
        self.assertEqual(extension, "xlsx")
        self.assertTrue(report_name)
        workbook = load_workbook(BytesIO(content))
        self.assertTrue(any(15 in row for row in workbook.active.values))

    def test_native_foreground_arabic_and_english_direction(self):
        name = self.owned_query()
        previous = (frappe.local.form_dict, frappe.local.response, frappe.local.lang)
        try:
            for language, expected in (("ar", True), ("en", False)):
                frappe.local.form_dict = self.params(name)
                frappe.local.response = frappe._dict()
                frappe.local.lang = language
                report_compat.export_query()
                workbook = load_workbook(BytesIO(frappe.local.response.filecontent))
                self.assertTrue(any(15 in row for row in workbook.active.values))
                self.assertEqual(bool(workbook.active.sheet_view.rightToLeft), expected)
        finally:
            frappe.local.form_dict, frappe.local.response, frappe.local.lang = previous

    def test_background_export_uses_native_data_arabic_and_owned_delivery_sink(self):
        # Exercise actual native export; capture delivery locally, never send email.
        name = self.owned_query()
        with patch("frappe.desk.utils.send_report_email") as delivery:
            report_compat.run_export_query_job("report-owned@example.invalid", self.params(name), frappe._dict(), "ar")
        delivery.assert_called_once()
        args, kwargs = delivery.call_args
        self.assertEqual(args[0], "report-owned@example.invalid")
        self.assertEqual(args[2], "xlsx")
        self.assertEqual(kwargs["attached_to_name"], name)
        workbook = load_workbook(BytesIO(args[3]))
        self.assertTrue(workbook.active.sheet_view.rightToLeft)
        self.assertTrue(any(15 in row for row in workbook.active.values))

    def test_serial_no_exact_native_empty_result_and_report_boundary(self):
        from erpnext.stock.report.available_serial_no import available_serial_no
        from frappe.core.doctype.report.report import Report
        data = available_serial_no.process_stock_ledger_entries(
            [frappe._dict(serial_and_batch_bundle=None)], {}, None, 2)
        self.assertIsInstance(data, tuple)
        self.assertEqual(data, ([], []))
        report = frappe.get_doc("Report", "Available Serial No")
        columns = [{"fieldname": "serial_no", "label": "Serial No", "fieldtype": "Data"}]
        with patch.object(Report, "execute_script_report", return_value=(columns, data)):
            result = report.execute_script_report(frappe._dict())
        self.assertIs(result[0], columns)
        self.assertEqual(result[1], [])

    def test_review_custom_query_preserved_exact_native_repair_transactional(self):
        upstream = json.loads(Path(frappe.get_app_path(
            "erpnext", "quality_management", "report", "review", "review.json")).read_text())
        meta = frappe.get_meta("Quality Action")
        self.assertTrue(meta.has_field("review"))
        self.assertFalse(meta.has_field("document_type"))
        self.assertIn(report_compat.OBSOLETE_REVIEW_COLUMN, upstream["query"])
        custom = upstream["query"] + "\n-- owned customization acceptance"
        frappe.db.set_value("Report", "Review", "query", custom, update_modified=False)
        frappe.clear_document_cache("Report", "Review")
        with patch.object(frappe.db, "commit", side_effect=AssertionError("internal commit")):
            self.assertFalse(report_compat.sync_report_compatibility()["review_query_updated"])
            self.assertEqual(frappe.db.get_value("Report", "Review", "query"), custom)
            frappe.db.set_value("Report", "Review", "query", upstream["query"], update_modified=False)
            frappe.clear_document_cache("Report", "Review")
            self.assertTrue(report_compat.sync_report_compatibility()["review_query_updated"])
            self.assertFalse(report_compat.sync_report_compatibility()["review_query_updated"])
        self.assertEqual(frappe.db.get_value("Report", "Review", "query"), report_compat.REVIEW_QUERY)

    def test_native_payment_permission_denial_adds_no_private_origin(self):
        email = "report-" + uuid.uuid4().hex + "@example.invalid"
        frappe.get_doc(dict(doctype="User", email=email, first_name="Report acceptance",
                            user_type="System User", send_welcome_email=0,
                            roles=[dict(role="Website Manager")])).insert(ignore_permissions=True)
        frappe.set_user(email)
        self.assertFalse(frappe.has_permission("Payment Entry", ptype="read"))
        result = {"columns": [{"fieldname": "voucher_no"}], "result": [
            {"voucher_type": "Payment Entry", "voucher_no": "BND nonexistent owned voucher"}]}
        native_get_list = frappe.get_list
        def permission_bounded_list(doctype, *args, **kwargs):
            if doctype == "Payment Entry":
                raise AssertionError("unauthorized payment lookup")
            return native_get_list(doctype, *args, **kwargs)
        with patch.object(frappe, "get_list", side_effect=permission_bounded_list):
            report_compat._add_ledger_statement(result)
        self.assertEqual(result["result"][0]["bnd_statement"], frappe._("Payment Entry"))

    def test_native_ledger_boundary_and_xlsx_serializer_preserve_owned_values(self):
        from frappe.core.doctype.report.report import Report
        from frappe.desk.query_report import build_xlsx_data
        from frappe.utils.xlsxutils import make_xlsx
        name = "BND PE report fixture " + uuid.uuid4().hex[:12]
        # Query-only owned row. No validation/submission/GL; this tests presentation.
        frappe.get_doc(dict(doctype="Payment Entry", name=name, payment_type="Receive",
                            total_allocated_amount=0)).db_insert()
        report = frappe.get_doc("Report", "General Ledger")
        columns = [frappe._dict(fieldname="voucher_no", label="Voucher No", fieldtype="Data", width=180),
                   frappe._dict(fieldname="voucher_type", label="Voucher Type", fieldtype="Data", width=180),
                   frappe._dict(fieldname="debit", label="Debit", fieldtype="Float", width=100)]
        rows = [frappe._dict(voucher_no=name, voucher_type="Payment Entry", debit=15)]
        # Isolate the measured native Report boundary, not financial query execution.
        with patch.object(Report, "execute_module", return_value=(columns, rows)):
            result = report.execute_script_report(frappe._dict())
        self.assertEqual(result[1][0]["debit"], 15)
        self.assertIn("bnd_statement", result[1][0])
        self.assertIn(result[1][0]["bnd_statement"],
                      (frappe._("Receipt Voucher"), frappe._("Receipt Voucher (origin not recorded)")))
        data = frappe._dict(columns=[frappe._dict(c) for c in result[0]], result=result[1],
                            report_name="General Ledger", filters={}, applied_filters={})
        table, widths, styles = build_xlsx_data(data, ignore_visible_idx=True)
        content = make_xlsx(table, "General Ledger", column_widths=widths, styles=styles).getvalue()
        values = list(load_workbook(BytesIO(content)).active.values)
        self.assertTrue(any(15 in row for row in values))
        self.assertTrue(any(result[1][0]["bnd_statement"] in row for row in values))

    def test_pdf_stylesheet_compatibility_keeps_native_renderer_and_boundaries(self):
        from bunood_theme.printing.report_assets import canonical_print_stylesheet
        from frappe.utils.jinja_globals import bundled_asset
        paths = {bundled_asset("print.bundle.css", rtl=rtl) for rtl in (False, True)}
        for path in paths:
            html = f'<link rel="stylesheet" href="http://browser-owned.localhost:8187{path}"><p dir="rtl">15</p>'
            result = canonical_print_stylesheet(html, "http://browser-owned.localhost:8187/",
                                                "http://team-rc.localhost/", paths)
            self.assertIn("http://team-rc.localhost" + path, result)
            self.assertIn('dir="rtl"', result)
            unrelated = '<link rel="stylesheet" href="http://browser-owned.localhost:8187/private.css">'
            self.assertEqual(canonical_print_stylesheet(unrelated, "http://browser-owned.localhost:8187/",
                                                        "http://team-rc.localhost/", paths), unrelated)
